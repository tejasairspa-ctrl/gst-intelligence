"""
GSTR-2A Compiler
────────────────
Merges N monthly GSTR-2A Excel files (portal downloads) into one consolidated
workbook that looks exactly like the GST portal format, with exactly two extra
columns prepended to every data sheet:

  Col A — GSTIN   : taxpayer's GSTIN (from Read me sheet)
  Col B — Period  : month-year label  e.g. "April 2024"

All other columns, header text, and styling replicate the portal exactly.
Rows are sorted in Financial-Year order (Apr first, Mar last).

Structural differences vs GSTR-2B (handled here):
  • Data sheets start right after the single "Read me" sheet (index 1+), not
    index 5 — a 2A file is [Read me, B2B, B2BA, CDNR, CDNRA, ECO, ECOA, ISD,
    ISDA, TDS, TDSA, TCS, IMPG, IMPG SEZ].
  • The Read me sheet stores GSTIN / Tax period / Financial year as label→value
    pairs in columns B→C and D→E; the tax period is "MMYYYY" (e.g. "042024").
  • Header block: rows 1-3 title, row 4 subtitle, rows 5-6 (regular) or 5-7
    (amendment) column headers; data starts row 7 (regular) / 8 (amendment) —
    same as 2B, so `_is_amendment_sheet` drives `data_start`.

Styling and the GSTIN/Period prepend logic are shared with the 2B compiler.
"""
import re
import io
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill

from services.gstr2b_compiler import (
    _MONTH_FULL, _fy_sort_key, _is_amendment_sheet,
    _apply_header_cell, _col_letter,
    _FILL_NAVY, _FILL_YELLOW, _FONT_HDR, _FONT_SUB, _FONT_DATA,
    _ALIGN_LEFT, _BORDER_DATA, _NAVY,
)


def _extract_readme_info(ws) -> tuple:
    """Return (period_label, gstin) from the GSTR-2A 'Read me' worksheet.

    Labels sit in column B (or D) with their value one cell to the right, e.g.
      B2="Taxpayer's GSTIN"  C2="19AAACN9225B1ZN"
      D2="Tax period"        E2="042024"
      D3="Financial year"    E3="2024-25"
    """
    fields: dict = {}
    for row in ws.iter_rows(min_row=1, max_row=15, values_only=True):
        for li, vi in ((1, 2), (3, 4)):
            if len(row) > vi and row[li] and row[vi] not in (None, ''):
                key = str(row[li]).strip().lower().rstrip(':')
                fields.setdefault(key, str(row[vi]).strip())

    gstin  = fields.get("taxpayer's gstin") or fields.get("gstin") or ''
    period = fields.get("tax period") or ''

    period_label = ''
    m = re.match(r'^(\d{2})(\d{4})$', period)
    if m:
        mn = int(m.group(1))
        if 1 <= mn <= 12:
            period_label = f"{_MONTH_FULL[mn - 1].capitalize()} {m.group(2)}"
    return period_label, gstin


def _period_from_filename(filename: str) -> str:
    """GSTR-2A portal filename: '<GSTIN>_<MMYYYY>_R2A.xlsx'."""
    m = re.search(r'_(\d{2})(\d{4})_', filename)
    if m:
        mn = int(m.group(1))
        if 1 <= mn <= 12:
            return f"{_MONTH_FULL[mn - 1].capitalize()} {m.group(2)}"
    return ''


# ── Main compiler ─────────────────────────────────────────────────────────────

def compile_gstr2a(file_streams: list) -> bytes:
    """
    Compile multiple GSTR-2A Excel files into one consolidated workbook.

    Parameters
    ----------
    file_streams : list of (filename: str, stream: BinaryIO)

    Returns
    -------
    bytes : compiled .xlsx workbook
    """
    registry:   dict = {}   # sheet_name → {header_rows, rows, max_col, data_start}
    files_meta: list = []

    for filename, stream in file_streams:
        try:
            wb = load_workbook(stream, data_only=True)
        except Exception as exc:
            raise ValueError(f"Cannot open '{filename}': {exc}") from exc

        period, gstin = _extract_readme_info(wb.worksheets[0])
        if not period:
            period = _period_from_filename(filename)
        if not period:
            period = filename[:10]

        files_meta.append({'filename': filename, 'gstin': gstin, 'period': period})

        # Data sheets = everything after the single "Read me" sheet.
        for ws in wb.worksheets[1:]:
            sname      = ws.title
            amendment  = _is_amendment_sheet(sname)
            data_start = 8 if amendment else 7

            if sname not in registry:
                header_rows = [
                    list(row)
                    for row in ws.iter_rows(min_row=1, max_row=data_start - 1, values_only=True)
                ]
                registry[sname] = {
                    'header_rows': header_rows,
                    'rows':        [],
                    'max_col':     ws.max_column or 21,
                    'data_start':  data_start,
                    'amendment':   amendment,
                }

            for row in ws.iter_rows(min_row=data_start, values_only=True):
                if any(c is not None for c in row):
                    registry[sname]['rows'].append((gstin, period, list(row)))

        wb.close()

    files_meta.sort(key=lambda fm: _fy_sort_key(fm['period']))

    # ── Build output workbook ─────────────────────────────────────────────
    wb_out = Workbook()
    wb_out.remove(wb_out.active)

    for sname, data in registry.items():
        ws_out   = wb_out.create_sheet(title=sname[:31])   # Excel 31-char cap
        hdr_rows = data['header_rows']
        n_hdr    = len(hdr_rows)
        max_col  = data['max_col']
        out_cols = max_col + 2

        for i, hrow in enumerate(hdr_rows):
            row_num = i + 1
            if row_num <= 4:
                out_row = [hrow[0] if hrow else None] + [None] * (out_cols - 1)
            elif row_num == n_hdr - 1:
                out_row = ['GSTIN', 'Period'] + hrow
            else:
                out_row = [None, None] + hrow
            ws_out.append(out_row)

            fill, font = (_FILL_YELLOW, _FONT_SUB) if row_num == 4 else (_FILL_NAVY, _FONT_HDR)
            for col_idx in range(1, out_cols + 1):
                _apply_header_cell(ws_out.cell(row=row_num, column=col_idx), fill, font)

        title_end = _col_letter(out_cols)
        ws_out.merge_cells(f'A1:{title_end}3')
        ws_out.merge_cells(f'A4:{title_end}4')
        first_hdr = n_hdr - 1
        ws_out.merge_cells(f'A{first_hdr}:A{n_hdr}')
        ws_out.merge_cells(f'B{first_hdr}:B{n_hdr}')

        sorted_rows = sorted(data['rows'], key=lambda r: _fy_sort_key(r[1]))
        for j, (gstin_val, period_val, src_row) in enumerate(sorted_rows):
            out_row = [gstin_val, period_val] + src_row
            ws_out.append(out_row)
            row_num = n_hdr + j + 1
            ws_out.row_dimensions[row_num].height = 15
            for col_idx in range(1, len(out_row) + 1):
                cell = ws_out.cell(row=row_num, column=col_idx)
                cell.font      = _FONT_DATA
                cell.alignment = _ALIGN_LEFT
                cell.border    = _BORDER_DATA

        ws_out.column_dimensions['A'].width = 22   # GSTIN
        ws_out.column_dimensions['B'].width = 14   # Period
        for col_letter, w in {'C': 20, 'D': 22, 'E': 16, 'F': 16, 'G': 18,
                              'H': 16, 'I': 14, 'J': 14, 'K': 14, 'L': 14}.items():
            ws_out.column_dimensions[col_letter].width = w

        ws_out.auto_filter.ref = f'A{n_hdr}:{title_end}{n_hdr}'
        ws_out.freeze_panes = f'C{n_hdr + 1}'

    # ── Compilation summary sheet ─────────────────────────────────────────
    info = wb_out.create_sheet(title='Summary', index=0)
    info.append(['GSTR-2A Compiled Report'])
    info.append([])
    info.append(['Period', 'GSTIN', 'Source File'])
    for fm in files_meta:
        info.append([fm['period'], fm['gstin'], fm['filename']])

    info['A1'].font = Font(name='Arial', bold=True, size=12, color=_NAVY)
    hdr_fill = PatternFill('solid', start_color=_NAVY)
    for col in range(1, 4):
        c = info.cell(row=3, column=col)
        c.fill = hdr_fill
        c.font = _FONT_HDR
    info.column_dimensions['A'].width = 18
    info.column_dimensions['B'].width = 22
    info.column_dimensions['C'].width = 55

    buf = io.BytesIO()
    wb_out.save(buf)
    buf.seek(0)
    return buf.getvalue()
